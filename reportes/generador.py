import os
from datetime import date, datetime
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


CARPETA_REPORTES = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "reportes_generados")
os.makedirs(CARPETA_REPORTES, exist_ok=True)


class GeneradorReportes:

    @staticmethod
    def _nombre_archivo(prefijo, extension):
        marca = datetime.now().strftime("%Y%m%d_%H%M%S")
        return os.path.join(CARPETA_REPORTES, f"{prefijo}_{marca}.{extension}")

    # ------------------------------------------------------------------
    # PDF
    # ------------------------------------------------------------------
    @classmethod
    def pdf_inventario(cls, bienes, titulo="Inventario general de bienes",
                       subtitulo="Departamento de Tesorería",
                       rango_fechas=None):
        ruta = cls._nombre_archivo("inventario", "pdf")
        doc = SimpleDocTemplate(ruta, pagesize=letter,
                                 leftMargin=2*cm, rightMargin=2*cm,
                                 topMargin=1.8*cm, bottomMargin=1.8*cm)
        estilos = getSampleStyleSheet()
        titulo_estilo = ParagraphStyle('titulo', parent=estilos['Title'],
                                        fontSize=16, textColor=colors.HexColor("#1a2b4a"))
        sub_estilo = ParagraphStyle('sub', parent=estilos['Normal'],
                                     fontSize=10, textColor=colors.HexColor("#6b7280"),
                                     alignment=1)

        contenido = [
            Paragraph("SAHUAYO — GOBIERNO", titulo_estilo),
            Paragraph(titulo, estilos['Heading2']),
            Paragraph(subtitulo, sub_estilo),
        ]
        if rango_fechas:
            contenido.append(Paragraph(f"Período: {rango_fechas}", sub_estilo))
        contenido.append(Paragraph(f"Fecha de generación: {datetime.now().strftime('%d/%m/%Y %H:%M')}", sub_estilo))
        contenido.append(Spacer(1, 0.5*cm))

        encabezados = ["No. Inventario", "Descripción", "Tipo", "Área", "Responsable", "Estado"]
        datos = [encabezados]
        for b in bienes:
            datos.append([
                b.get("codigo_inventario", ""),
                b.get("nombre", "")[:40],
                b.get("tipo", "")[:20],
                b.get("area", "")[:20],
                b.get("responsable", "")[:25],
                b.get("estado", ""),
            ])

        tabla = Table(datos, colWidths=[2.5*cm, 5*cm, 3*cm, 3*cm, 4*cm, 2*cm], repeatRows=1)
        tabla.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#dbeafe")),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.HexColor("#1a2b4a")),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('GRID', (0, 0), (-1, -1), 0.3, colors.HexColor("#e5e7eb")),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor("#fafbfc")]),
            ('TOPPADDING', (0, 0), (-1, -1), 5),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ]))
        contenido.append(tabla)
        contenido.append(Spacer(1, 0.5*cm))
        contenido.append(Paragraph(f"Total de registros: {len(bienes)}", sub_estilo))

        doc.build(contenido)
        return ruta

    # ------------------------------------------------------------------
    # Excel
    # ------------------------------------------------------------------
    @classmethod
    def excel_inventario(cls, bienes, titulo="Inventario general de bienes"):
        ruta = cls._nombre_archivo("inventario", "xlsx")
        wb = Workbook()
        ws = wb.active
        ws.title = "Inventario"

        ws.merge_cells("A1:F1")
        ws["A1"] = "SAHUAYO — GOBIERNO"
        ws["A1"].font = Font(bold=True, size=14, color="1a2b4a")
        ws["A1"].alignment = Alignment(horizontal="center")

        ws.merge_cells("A2:F2")
        ws["A2"] = titulo
        ws["A2"].font = Font(bold=True, size=12)
        ws["A2"].alignment = Alignment(horizontal="center")

        ws.merge_cells("A3:F3")
        ws["A3"] = f"Fecha de generación: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
        ws["A3"].font = Font(size=9, italic=True, color="6b7280")
        ws["A3"].alignment = Alignment(horizontal="center")

        encabezados = ["No. Inventario", "Descripción", "Tipo", "Área", "Responsable", "Estado"]
        fila_enc = 5
        for i, h in enumerate(encabezados, start=1):
            c = ws.cell(row=fila_enc, column=i, value=h)
            c.font = Font(bold=True, color="1a2b4a")
            c.fill = PatternFill("solid", fgColor="dbeafe")
            c.alignment = Alignment(horizontal="center")
            c.border = Border(bottom=Side(style="thin", color="e5e7eb"))

        for idx, b in enumerate(bienes, start=fila_enc + 1):
            ws.cell(row=idx, column=1, value=b.get("codigo_inventario", ""))
            ws.cell(row=idx, column=2, value=b.get("nombre", ""))
            ws.cell(row=idx, column=3, value=b.get("tipo", ""))
            ws.cell(row=idx, column=4, value=b.get("area", ""))
            ws.cell(row=idx, column=5, value=b.get("responsable", ""))
            ws.cell(row=idx, column=6, value=b.get("estado", ""))

        anchos = [16, 40, 20, 20, 25, 12]
        for i, ancho in enumerate(anchos, start=1):
            ws.column_dimensions[get_column_letter(i)].width = ancho

        wb.save(ruta)
        return ruta